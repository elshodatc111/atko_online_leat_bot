"""Excel eksport — fayl tayyorlanib, so'ragan xodimga bot orqali yuboriladi."""
from __future__ import annotations

import html
from datetime import date
from pathlib import Path

from aiogram.types import FSInputFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from ..config import config
from ..db import session_scope
from ..models import FORMATS, GOALS, TEMPERATURES, AuditLog, Chat, Lead, Payment, Subscription, UserQuestion
from . import audit as audit_svc
from . import stats, worktime

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _out_path(name: str) -> Path:
    folder = config.DATA_DIR / "exports"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"{name}_{worktime.now_local().strftime('%Y%m%d_%H%M%S')}.xlsx"


async def leads_xlsx(start: date | None = None, end: date | None = None, status: str | None = None) -> Path:
    async with session_scope() as s:
        q = select(Lead).order_by(Lead.id.desc())
        if start and end:
            s0, s1 = stats.range_utc(start, end)
            q = q.where(Lead.created_at >= s0, Lead.created_at < s1)
        if status:
            q = q.where(Lead.status == status)
        leads = (await s.execute(q)).scalars().all()
    rows = [[
        l.id, l.name, l.phone, f"@{l.tg_username}" if l.tg_username else "", l.tg_id, l.lang,
        GOALS.get(l.goal or "", {}).get("uz", l.goal), FORMATS.get(l.study_format or "", {}).get("uz", l.study_format),
        l.level, l.city, l.interested_tariff, l.source.name if l.source else "", l.status_label, l.reject_reason,
        TEMPERATURES.get(l.temperature or "", ""), l.last_operator.display_name if l.last_operator else "",
        l.accepted_by.display_name if l.accepted_by else "", worktime.fmt(l.accepted_at) if l.accepted_at else "",
        "Ha" if l.is_blocked else "", worktime.fmt(l.created_at),
        worktime.fmt(l.last_activity), l.ai_summary,
    ] for l in leads]
    wb = Workbook()
    _sheet(wb, "Leadlar", ["ID", "Ism", "Telefon", "Username", "Telegram ID", "Til", "Maqsad", "Format", "Daraja",
                           "Shahar", "Qiziqqan tarif", "Manba", "Status", "Rad etish sababi", "Qiziqish",
                           "Oxirgi operator", "Qabul qilgan", "Qabul sanasi", "Botni bloklagan",
                           "Yaratilgan", "Oxirgi faollik", "AI xulosa"], rows, first=True)
    p = _out_path("leadlar")
    wb.save(p)
    return p


async def chats_xlsx(start: date, end: date) -> Path:
    s0, s1 = stats.range_utc(start, end)
    async with session_scope() as s:
        chats = (await s.execute(select(Chat).where(Chat.requested_at >= s0, Chat.requested_at < s1)
                                 .order_by(Chat.id.desc()))).scalars().all()
    def sec(a, b):
        return stats.human_seconds((a - b).total_seconds()) if a and b else ""
    rows = [[
        c.id, c.lead.display if c.lead else "", c.lead.phone if c.lead else "", c.status_label,
        c.operator.display_name if c.operator else "", "Ha" if c.off_hours else "", worktime.fmt(c.requested_at),
        worktime.fmt(c.accepted_at) if c.accepted_at else "", sec(c.accepted_at, c.sla_start),
        sec(c.first_response_at, c.accepted_at), worktime.fmt(c.closed_at) if c.closed_at else "",
        sec(c.closed_at, c.accepted_at), c.rating or "", c.rating_comment or "", c.transfers,
    ] for c in chats]
    wb = Workbook()
    _sheet(wb, "Chatlar", ["ID", "Lead", "Telefon", "Holat", "Operator", "Ish vaqtidan tashqari", "So'rov vaqti",
                           "Olingan vaqt", "Kutish", "Birinchi javob", "Yopilgan", "Davomiylik", "Baho", "Izoh",
                           "O'tkazishlar"], rows, first=True)
    p = _out_path("chatlar")
    wb.save(p)
    return p


async def operators_xlsx(start: date, end: date) -> Path:
    data = await stats.operator_stats(start, end)
    rows = [[d["name"], d["display"], "Admin" if d["role"] == "admin" else "Operator", "Faol" if d["active"] else "O'chirilgan",
             d["chats"], d["closed"], d["open"], stats.human_seconds(d["avg_wait"]),
             stats.human_seconds(d["avg_first_response"]), stats.human_seconds(d["avg_duration"]),
             d["avg_rating"] or "", d["ratings"], d["low_ratings"], d["accepted"], d["conv"], d["messages"],
             d["status_changes"], d["transfers"]] for d in data]
    wb = Workbook()
    _sheet(wb, "Operatorlar", ["F.I.Sh", "Leadga ko'rinadigan ism", "Rol", "Holat", "Olingan chatlar", "Yopilgan",
                               "Ochiq", "O'rtacha kutish", "O'rtacha birinchi javob", "O'rtacha davomiylik",
                               "O'rtacha baho", "Baholar soni", "Past baholar (1-2)", "Kursga qabul qilingan",
                               "Konversiya %", "Yuborilgan xabarlar", "Status o'zgarishlari", "O'tkazishlar"], rows, first=True)
    # kunlik kesim
    series = await stats.operator_series(start, end)
    daily_rows = []
    for i, label in enumerate(series["labels"]):
        daily_rows.append([label] + [ds["data"][i] for ds in series["datasets"]])
    _sheet(wb, "Kunlik chatlar", ["Sana"] + [ds["label"] for ds in series["datasets"]], daily_rows)
    p = _out_path("operatorlar")
    wb.save(p)
    return p


async def audit_xlsx(start: date, end: date) -> Path:
    s0, s1 = stats.range_utc(start, end)
    async with session_scope() as s:
        rows_db = (await s.execute(select(AuditLog).where(AuditLog.created_at >= s0, AuditLog.created_at < s1)
                                   .order_by(AuditLog.id.desc()))).scalars().all()
    rows = [[r.id, worktime.fmt(r.created_at), r.staff.full_name if r.staff else "Tizim",
             audit_svc.ACTIONS.get(r.action, r.action), r.entity or "", r.entity_id or "", r.details or ""] for r in rows_db]
    wb = Workbook()
    _sheet(wb, "Harakatlar jurnali", ["ID", "Vaqt", "Xodim", "Harakat", "Obyekt", "Obyekt ID", "Tafsilot"], rows, first=True)
    p = _out_path("harakatlar_jurnali")
    wb.save(p)
    return p


async def questions_xlsx(start: date, end: date) -> Path:
    s0, s1 = stats.range_utc(start, end)
    async with session_scope() as s:
        rows_db = (await s.execute(select(UserQuestion).where(UserQuestion.created_at >= s0, UserQuestion.created_at < s1)
                                   .order_by(UserQuestion.id.desc()))).scalars().all()
    rows = [[q.id, worktime.fmt(q.created_at), q.lead.display if q.lead else "", q.lang, q.mode, q.text,
             "Ha" if q.answered else "Yo'q"] for q in rows_db]
    wb = Workbook()
    _sheet(wb, "Savollar", ["ID", "Vaqt", "Lead", "Til", "Rejim", "Savol", "AI javob topdi"], rows, first=True)
    p = _out_path("savollar")
    wb.save(p)
    return p


async def payments_xlsx(start: date, end: date) -> Path:
    s0, s1 = stats.range_utc(start, end)
    async with session_scope() as s:
        rows_db = (await s.execute(select(Payment).where(Payment.created_at >= s0, Payment.created_at < s1)
                                   .order_by(Payment.id.desc()))).scalars().all()
    rows = [[p.id, worktime.fmt(p.created_at), worktime.fmt(p.paid_at) if p.paid_at else "", p.lead.display if p.lead else "",
             p.lead.phone if p.lead else "", p.tg_id, p.title, p.days, p.amount, p.state_label, p.payme_id or "",
             "Ha" if p.is_test else ""] for p in rows_db]
    wb = Workbook()
    _sheet(wb, "To'lovlar", ["№", "Yaratilgan", "To'langan", "Foydalanuvchi", "Telefon", "Telegram ID", "Tarif", "Kun",
                             "Summa (so'm)", "Holat", "Payme ID", "Test"], rows, first=True)
    p = _out_path("tolovlar")
    wb.save(p)
    return p


async def subscriptions_xlsx(start: date | None = None, end: date | None = None) -> Path:
    async with session_scope() as s:
        rows_db = (await s.execute(select(Subscription).order_by(Subscription.expires_at))).scalars().all()
    rows = [[x.id, x.lead.display if x.lead else (x.name or ""), x.lead.phone if x.lead else "", x.tg_id,
             worktime.fmt(x.started_at), "muddatsiz" if x.whitelisted else worktime.fmt(x.expires_at),
             "∞" if x.whitelisted else x.days_left, "Faol" if x.is_active else "Tugagan", "Ha" if x.in_group else "",
             x.note or ""] for x in rows_db]
    wb = Workbook()
    _sheet(wb, "Obunachilar", ["ID", "Foydalanuvchi", "Telefon", "Telegram ID", "Boshlangan", "Tugaydi", "Qoldi (kun)",
                               "Holat", "Guruhda", "Izoh"], rows, first=True)
    p = _out_path("obunachilar")
    wb.save(p)
    return p


async def send_to_staff(tg_id: int, path: Path, caption: str) -> None:
    from ..bot.instance import get_bot

    await get_bot().send_document(tg_id, FSInputFile(path, filename=path.name), caption=html.escape(caption))
